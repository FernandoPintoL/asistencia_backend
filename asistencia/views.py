"""
API JSON que consume el frontend (React). La lógica de negocio sigue en services.py;
aquí solo se reciben datos, se llama a los servicios y se responde en JSON.
"""
from io import BytesIO

from django.contrib.auth import authenticate
from django.core.paginator import Paginator
from django.db.models import Count, Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from rest_framework import status
from rest_framework.decorators import api_view, authentication_classes, permission_classes
from rest_framework.response import Response

from .auth import EsAdmin, EsEstudiante, token_admin, token_estudiante
from .importers import parse_estudiantes, parse_fechas
from .models import Actividad, Asistencia, Estudiante, FechaClase, Modulo
from .services import (ETIQUETAS, activar_modulo, deshabilitar_modulo, eliminar_modulo, estado_calculado,
                       habilitar_modulo, marcar_asistencia, modificar_asistencia)

SIN_MODULO = "No hay un módulo activo. Habilita uno existente o importa uno nuevo."

# ───────────────────────── Conversión a JSON ─────────────────────────


def _momento(dt):
    """Fecha y hora en la zona de Bolivia, en formato ISO."""
    return timezone.localtime(dt).isoformat() if dt else None


def _modulo_json(m):
    return {"id": m.id, "nombre": m.nombre} if m else None


def _fecha_json(f):
    return {"id": f.id, "fecha": f.fecha.isoformat(),
            "inicio": f.hora_inicio.strftime("%H:%M"), "fin": f.hora_fin.strftime("%H:%M")}


def _estudiante_json(e):
    return {"id": e.id, "nro_registro": e.nro_registro, "nombre": e.nombre, "correo": e.correo}


def _actividad_json(a):
    return {"id": a.id, "fecha_hora": _momento(a.fecha_hora), "usuario": a.usuario, "rol": a.rol,
            "accion": a.accion, "registro_afectado": a.registro_afectado,
            "valor_anterior": a.valor_anterior, "valor_nuevo": a.valor_nuevo}


def _paginar(request, items, a_json):
    page = Paginator(items, 25).get_page(request.query_params.get("page"))
    return {"count": page.paginator.count, "page": page.number, "num_pages": page.paginator.num_pages,
            "results": [a_json(x) for x in page]}


def _sin_modulo():
    return Response({"detail": SIN_MODULO, "sin_modulo": True}, status=status.HTTP_409_CONFLICT)


def _error(msg, codigo=status.HTTP_400_BAD_REQUEST):
    return Response({"detail": msg}, status=codigo)


# ───────────────────────── Público ─────────────────────────


@api_view(["GET"])
@authentication_classes([])
def modulo_activo(request):
    return Response({"modulo": _modulo_json(Modulo.activo_actual())})


# ───────────────────────── Estudiante ─────────────────────────


@api_view(["POST"])
@authentication_classes([])
def login_estudiante(request):
    modulo = Modulo.activo_actual()
    dato = str(request.data.get("identificador", "")).strip()
    est = None
    if modulo and dato:
        est = (Estudiante.objects.filter(modulo=modulo)
               .filter(Q(correo__iexact=dato) | Q(nro_registro=dato)).first())
    if est is None:
        return _error("Esos datos no corresponden a un estudiante del módulo activo. "
                      "Revisa tu correo o número de registro.")
    return Response({"token": token_estudiante(est), "nombre": est.nombre})


@api_view(["GET"])
@permission_classes([EsEstudiante])
def mi_asistencia(request):
    est = request.user.estudiante
    hoy = timezone.localdate()
    registros = {a.fecha_clase_id: a for a in est.asistencias.all()}
    filas, presentes, faltas = [], 0, 0
    clase_hoy = None
    for f in est.modulo.fechas.all():
        reg = registros.get(f.id)
        estado = estado_calculado(f, reg, hoy)
        presentes += estado == "PRESENTE"
        faltas += estado == "FALTA"
        if f.fecha == hoy:
            clase_hoy = {"fecha": _fecha_json(f), "registrada": reg is not None}
        filas.append({
            "fecha": _fecha_json(f), "estado": estado, "es_hoy": f.fecha == hoy,
            "hora_marca": _momento(reg.marcado_en) if reg else None,
            "corregido": bool(reg and reg.modificado_por_admin),
        })
    return Response({
        "estudiante": {**_estudiante_json(est), "modulo": est.modulo.nombre},
        "filas": filas, "presentes": presentes, "faltas": faltas, "clase_hoy": clase_hoy,
        "puede_marcar": bool(clase_hoy and not clase_hoy["registrada"]),
    })


@api_view(["POST"])
@permission_classes([EsEstudiante])
def marcar(request):
    ok, msg = marcar_asistencia(request.user.estudiante)
    return Response({"detail": msg}) if ok else _error(msg)


# ───────────────────────── Administrador ─────────────────────────


@api_view(["POST"])
@authentication_classes([])
def login_admin(request):
    user = authenticate(request, username=request.data.get("username", ""),
                        password=request.data.get("password", ""))
    if user is None or not user.is_staff:
        return _error("Usuario o contraseña incorrectos, o sin permisos de administrador.")
    return Response({"token": token_admin(user), "nombre": user.get_username()})


@api_view(["GET"])
@permission_classes([EsAdmin])
def dashboard(request):
    modulo = Modulo.activo_actual()
    if not modulo:
        return Response({"modulo": None})
    hoy = timezone.localdate()
    clase_hoy = modulo.fechas.filter(fecha=hoy).first()
    proxima = modulo.fechas.filter(fecha__gt=hoy).first()
    return Response({
        "modulo": _modulo_json(modulo),
        "total_estudiantes": modulo.estudiantes.count(),
        "total_fechas": modulo.fechas.count(),
        "hay_clase_hoy": clase_hoy is not None,
        "presentes_hoy": (Asistencia.objects.filter(fecha_clase=clase_hoy, estado="PRESENTE").count()
                          if clase_hoy else 0),
        "proxima": proxima.fecha.isoformat() if proxima else None,
        "ultimas": [_actividad_json(a) for a in Actividad.objects.filter(modulo=modulo)[:6]],
    })


@api_view(["GET"])
@permission_classes([EsAdmin])
def estudiantes(request):
    modulo = Modulo.activo_actual()
    if not modulo:
        return _sin_modulo()
    q = request.query_params.get("q", "").strip()
    qs = modulo.estudiantes.all()
    if q:
        qs = qs.filter(Q(nombre__icontains=q) | Q(nro_registro__icontains=q))
    return Response({"modulo": _modulo_json(modulo), **_paginar(request, qs, _estudiante_json)})


def _filas_asistencia(modulo, q, fecha_id):
    hoy = timezone.localdate()
    est_qs = modulo.estudiantes.all()
    if q:
        est_qs = est_qs.filter(Q(nombre__icontains=q) | Q(nro_registro__icontains=q))
    fechas = list(modulo.fechas.all())
    if fecha_id:
        fechas = [f for f in fechas if str(f.id) == fecha_id]
    estudiantes_l = list(est_qs)
    regs = {(a.estudiante_id, a.fecha_clase_id): a
            for a in Asistencia.objects.filter(estudiante__in=est_qs, fecha_clase__in=[f.id for f in fechas])}
    filas = []
    for f in fechas:
        for e in estudiantes_l:
            reg = regs.get((e.id, f.id))
            filas.append({
                "estudiante": _estudiante_json(e), "fecha": _fecha_json(f),
                "estado": estado_calculado(f, reg, hoy),
                "hora_marca": _momento(reg.marcado_en) if reg else None,
                "corregido": bool(reg and reg.modificado_por_admin),
            })
    return filas


@api_view(["GET"])
@permission_classes([EsAdmin])
def asistencias(request):
    modulo = Modulo.activo_actual()
    if not modulo:
        return _sin_modulo()
    q = request.query_params.get("q", "").strip()
    fecha_id = request.query_params.get("fecha", "")
    filas = _filas_asistencia(modulo, q, fecha_id)
    totales = {
        "presentes": sum(f["estado"] == "PRESENTE" for f in filas),
        "faltas": sum(f["estado"] == "FALTA" for f in filas),
        "pendientes": sum(f["estado"] in ("PENDIENTE", "PROGRAMADA") for f in filas),
    }
    return Response({
        "modulo": _modulo_json(modulo), "fechas": [_fecha_json(f) for f in modulo.fechas.all()],
        "totales": totales, **_paginar(request, filas, lambda f: f),
    })


def _id(valor):
    try:
        return int(valor)
    except (TypeError, ValueError):
        return None


@api_view(["POST"])
@permission_classes([EsAdmin])
def modificar(request):
    modulo = Modulo.activo_actual()
    est = get_object_or_404(Estudiante, pk=_id(request.data.get("estudiante")), modulo=modulo)
    fecha = get_object_or_404(FechaClase, pk=_id(request.data.get("fecha")), modulo=modulo)
    ok, msg = modificar_asistencia(request.user, est, fecha, request.data.get("estado"))
    return Response({"detail": msg}) if ok else _error(msg)


@api_view(["GET"])
@permission_classes([EsAdmin])
def actividades(request):
    modulo = Modulo.activo_actual()
    if not modulo:
        return _sin_modulo()
    q = request.query_params.get("q", "").strip()
    fecha = request.query_params.get("fecha", "").strip()
    qs = Actividad.objects.filter(modulo=modulo)
    if q:
        qs = qs.filter(Q(usuario__icontains=q) | Q(registro_afectado__icontains=q))
    if fecha:
        qs = qs.filter(fecha_hora__date=fecha)
    return Response({"modulo": _modulo_json(modulo), **_paginar(request, qs, _actividad_json)})


def _leer_importacion(request):
    """Valida nombre y los dos Excel. Devuelve (nombre, estudiantes, fechas, errores)."""
    nombre = str(request.data.get("nombre", "")).strip()
    f_est, f_fec = request.FILES.get("estudiantes"), request.FILES.get("fechas")
    errores = []
    if not nombre:
        errores.append("Indica un nombre para el nuevo módulo.")
    if not f_est:
        errores.append("Falta el Excel de estudiantes.")
    if not f_fec:
        errores.append("Falta el Excel de fechas.")
    if errores:
        return nombre, [], [], errores
    est, e1 = parse_estudiantes(f_est)
    fec, e2 = parse_fechas(f_fec)
    errores = [f"Estudiantes · {e}" for e in e1] + [f"Fechas · {e}" for e in e2]
    return nombre, est, fec, errores


@api_view(["POST"])
@permission_classes([EsAdmin])
def modulo_validar(request):
    """RF-16: vista previa antes de activar. No guarda nada."""
    nombre, est, fec, errores = _leer_importacion(request)
    if errores:
        return Response({"errores": errores}, status=status.HTTP_400_BAD_REQUEST)
    return Response({"nombre": nombre, "estudiantes": est, "fechas": fec})


@api_view(["POST"])
@permission_classes([EsAdmin])
def modulo_confirmar(request):
    """
    El frontend vuelve a enviar los mismos archivos y aquí se validan de nuevo:
    así no hace falta guardar nada en sesión ni confiar en datos del navegador.
    """
    nombre, est, fec, errores = _leer_importacion(request)
    if errores:
        return Response({"errores": errores}, status=status.HTTP_400_BAD_REQUEST)
    modulo = activar_modulo(request.user, nombre, est, fec)
    return Response({"detail": f"Módulo «{modulo.nombre}» activado. Los datos del módulo anterior "
                               "se conservan en la base de datos."}, status=status.HTTP_201_CREATED)


@api_view(["GET"])
@permission_classes([EsAdmin])
def exportar(request):
    """RF-17: asistencias del módulo activo en Excel (estudiantes × fechas)."""
    modulo = Modulo.activo_actual()
    if not modulo:
        return _sin_modulo()
    hoy = timezone.localdate()
    fechas = list(modulo.fechas.all())
    regs = {(a.estudiante_id, a.fecha_clase_id): a
            for a in Asistencia.objects.filter(estudiante__modulo=modulo)}
    wb = Workbook()
    ws = wb.active
    ws.title = "Asistencias"
    ws.append(["Registro", "Nombre", "Correo"] + [f.fecha.strftime("%d/%m/%Y") for f in fechas]
              + ["Presentes", "Faltas"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="17233B")
    for e in modulo.estudiantes.all():
        estados = [estado_calculado(f, regs.get((e.id, f.id)), hoy) for f in fechas]
        ws.append([e.nro_registro, e.nombre, e.correo] + [ETIQUETAS[s] for s in estados]
                  + [estados.count("PRESENTE"), estados.count("FALTA")])
    for col, ancho in zip("ABC", (12, 30, 30)):
        ws.column_dimensions[col].width = ancho
    buffer = BytesIO()
    wb.save(buffer)
    resp = HttpResponse(buffer.getvalue(),
                        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="asistencias_{hoy:%Y%m%d}.xlsx"'
    return resp


@api_view(["GET"])
@permission_classes([EsAdmin])
def modulos(request):
    """Lista de todos los módulos, con opción de habilitar o deshabilitar."""
    lista = (Modulo.objects
             .annotate(n_estudiantes=Count("estudiantes", distinct=True),
                       n_fechas=Count("fechas", distinct=True),
                       n_asistencias=Count("estudiantes__asistencias", distinct=True))
             .order_by("-creado_en", "-id"))
    return Response([{
        "id": m.id, "nombre": m.nombre, "activo": m.activo, "creado_en": _momento(m.creado_en),
        "n_estudiantes": m.n_estudiantes, "n_fechas": m.n_fechas, "n_asistencias": m.n_asistencias,
    } for m in lista])


@api_view(["POST"])
@permission_classes([EsAdmin])
def modulo_habilitar(request, pk):
    modulo = get_object_or_404(Modulo, pk=pk)
    habilitar_modulo(request.user, modulo)
    return Response({"detail": f"Módulo «{modulo.nombre}» habilitado. Solo sus estudiantes pueden ingresar ahora."})


@api_view(["POST"])
@permission_classes([EsAdmin])
def modulo_deshabilitar(request, pk):
    modulo = get_object_or_404(Modulo, pk=pk)
    deshabilitar_modulo(request.user, modulo)
    return Response({"detail": f"Módulo «{modulo.nombre}» deshabilitado. Ningún estudiante puede ingresar "
                               "hasta habilitar un módulo."})


@api_view(["POST"])
@permission_classes([EsAdmin])
def modulo_eliminar(request, pk):
    modulo = get_object_or_404(Modulo, pk=pk)
    ok, msg = eliminar_modulo(request.user, modulo, request.data.get("confirmacion", ""))
    return Response({"detail": msg}) if ok else _error(msg)
